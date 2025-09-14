import logging
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
from utils.excel_helper import red_fill, yellow_fill, green_fill, orange_fill, blue_fill, pink_fill
from config import COLUMNS_FOR_EXCEL_FORMATTER
from logging_config import setup_logging

setup_logging()
logger = logging.getLogger(__name__)


class ExcelFormatter:
    def __init__(self, df):
        self.work_book = Workbook()
        self.df = df

    def get_excel_for_comparison(self):
        logger.info('Генерация excel файла')
        ws = self.work_book.active
        self._write_dataframe(ws)
        col_letters = self._map_columns()
        # self._add_new_columns(ws, col_letters)
        self._make_headers_bold(ws)
        self._fill_formulas(ws, col_letters)
        self._apply_conditional_formatting(ws, col_letters)
        self._auto_fit_columns(ws)
        self.work_book.save("compare_price.xlsx")

    def _write_dataframe(self, ws):
        for row in dataframe_to_rows(self.df, index=False, header=True):
            ws.append(row)

    def _map_columns(self) -> dict:
        return {col: get_column_letter(i + 1) for i, col in enumerate(self.df.columns)}

    def _make_headers_bold(self, ws):
        """Делает заголовки столбцов жирными"""
        bold_font = Font(bold=True)
        for cell in ws[1]:  # Первая строка содержит заголовки
            cell.font = bold_font

    def _fill_formulas(self, ws, cols):
        for row in range(2, len(self.df) + 2):

            ws[f"{cols[COLUMNS_FOR_EXCEL_FORMATTER['discount_wb']]}{row}"].fill = pink_fill
            ws[f"{cols[COLUMNS_FOR_EXCEL_FORMATTER['seller_discount']]}{row}"].fill = green_fill
            columns_for_blue = ['(WB) Цена со скидкой WB\n(черная)','(MED) Цена со скидкой продавца']
            for col_name in columns_for_blue:
                ws[f"{cols[col_name]}{row}"].fill = blue_fill

    def _apply_conditional_formatting(self, ws, cols):
        last_row = len(self.df) + 2
        price_range = f"{cols[COLUMNS_FOR_EXCEL_FORMATTER['price_diff']]}2:" \
                      f"{cols[COLUMNS_FOR_EXCEL_FORMATTER['price_diff']]}{last_row}"
        rule_red = FormulaRule(
            formula=[f"{cols[COLUMNS_FOR_EXCEL_FORMATTER['med_price']]}2 - "
                     f"{cols[COLUMNS_FOR_EXCEL_FORMATTER['wb_price']]}2 >= 1000"],
            stopIfTrue=True,
            fill=red_fill
        )

        rule_orange = FormulaRule(
            formula=[f"{cols[COLUMNS_FOR_EXCEL_FORMATTER['wb_price']]}2 - "
                     f"{cols[COLUMNS_FOR_EXCEL_FORMATTER['med_price']]}2 >= 1000"],
            stopIfTrue=True,
            fill=orange_fill
        )

        ws.conditional_formatting.add(price_range, rule_red)
        ws.conditional_formatting.add(price_range, rule_orange)

    @staticmethod
    def _auto_fit_columns(ws):
        for col in ws.columns:
            max_length = max((len(str(cell.value)) for cell in col if cell.value), default=0)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max_length + 2
