from openpyxl.styles import PatternFill
import xlwings as xw
import pandas as pd
from DTO.columns_dto import WBColumnsDTO
from DTO.wb_dop_columns import WBDopColumnsDTO

columns_dto = WBColumnsDTO()
dop_columns_dto = WBDopColumnsDTO()

yellow_fill = PatternFill(start_color='F8FCC4',  # Желтый цвет
                          end_color='F8FCC4',
                          fill_type='solid')
green_fill = PatternFill(start_color='C9EAA9',  # Зеленый цвет
                         end_color='C9EAA9',
                         fill_type='solid')
light_green_fill = PatternFill(start_color='60e755',  # Зеленый цвет
                         end_color='60e755',
                         fill_type='solid')

red_fill = PatternFill(start_color='FFC7CE',  # Красный цвет (Розовый цвет)
                       end_color='FFC7CE',
                       fill_type='solid')
orange_fill = PatternFill(start_color="FFEB9C", # Оранжевый цвет (Желтоватый цвет)
                          end_color="FFEB9C",
                          fill_type="solid")
blue_fill = PatternFill(start_color="9FD6E5",
                          end_color="9FD6E5",
                          fill_type="solid")
ozon_color_fill = PatternFill(start_color="5295CB",
                              end_color="5295CB",
                              fill_type="solid")

pink_fill = PatternFill(start_color="EF99D8",
                        end_color="EF99D8",
                        fill_type="solid")


def load_excel_with_formulas(filename: str) -> pd.DataFrame:
    app = xw.App(visible=False)  # Excel в фоне
    wb = app.books.open(filename)

    # Принудительно пересчитать все формулы
    wb.app.calculate()

    ws = wb.sheets[0]  # Первый лист
    df = ws.used_range.options(pd.DataFrame, header=1, index=False).value

    wb.close()
    app.quit()

    percent_columns = [dop_columns_dto.mp_commission,
                       dop_columns_dto.mu_original,
                       dop_columns_dto.mu_with_our_discount,
                       dop_columns_dto.mu_with_discount_from_the_price_with_spp,
                       dop_columns_dto.mu_taking_into_account_the_wb_commission,
                       dop_columns_dto.max_discount_including_commission]
    for col in percent_columns:
        if col in df.columns:
            df[col] = (df[col] * 100).round().astype(int)

    return df
